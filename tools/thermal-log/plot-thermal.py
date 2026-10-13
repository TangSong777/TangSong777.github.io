#!/usr/bin/env python3
"""把 thermal-log.mjs 采集的 CSV 转成带折线图与统计摘要的 xlsx。

用法：
    python plot-thermal.py 输入.csv [输出.xlsx]

生成两个工作表：
  data    原始数据（数值列已转成数字）
  report  统计摘要（各域温度 min/平均/max、超过各阈值的时长占比、风扇运行占比、SSD 温度）

并在 data 表右侧绘制图表：SoC 各域温度、风扇 PWM 与负载、CPU 频率、SSD 温度。
"""
import csv
import sys
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

TEXT_COLUMNS = {'iso', 'ssd_source'}
TEMPERATURE_COLUMNS = ['soc_c', 'bigcore0_c', 'bigcore1_c', 'littlecore_c', 'center_c', 'gpu_c', 'npu_c', 'max_c']
THRESHOLDS = [55, 60, 65, 70, 75, 80, 85]


def to_number(value):
    if value is None or value == '':
        return None
    try:
        return float(value)
    except ValueError:
        return None


def load_rows(path):
    with open(path, newline='', encoding='utf-8') as handle:
        return list(csv.DictReader(handle))


def numeric_column(rows, key):
    return [value for value in (to_number(row.get(key)) for row in rows) if value is not None]


def summarize(rows):
    lines = []
    lines.append(('采样行数', len(rows)))
    if rows:
        lines.append(('起始时间(ISO)', rows[0].get('iso', '')))
        lines.append(('结束时间(ISO)', rows[-1].get('iso', '')))
        span = (to_number(rows[-1].get('timestamp')) or 0) - (to_number(rows[0].get('timestamp')) or 0)
        lines.append(('覆盖时长(小时)', round(span / 3600, 2)))
    lines.append(('', ''))

    lines.append(('各热区温度', '最低 / 平均 / 最高 (°C)'))
    for key in TEMPERATURE_COLUMNS:
        values = numeric_column(rows, key)
        if not values:
            continue
        lines.append((key, f'{min(values):.1f} / {sum(values) / len(values):.1f} / {max(values):.1f}'))
    lines.append(('', ''))

    maxima = numeric_column(rows, 'max_c')
    if maxima:
        lines.append(('阈值统计（按 max_c）', '样本数 / 占比'))
        total = len(maxima)
        for threshold in THRESHOLDS:
            over = [value for value in maxima if value >= threshold]
            lines.append((f'≥ {threshold}°C', f'{len(over)} / {len(over) / total * 100:.1f}%'))
    lines.append(('', ''))

    pwm = numeric_column(rows, 'pwm')
    if pwm:
        running = [value for value in pwm if value > 0]
        lines.append(('风扇转速档位', '最低 / 平均 / 最高'))
        lines.append(('pwm', f'{min(pwm):.0f} / {sum(pwm) / len(pwm):.1f} / {max(pwm):.0f}'))
        lines.append(('风扇运行占比', f'{len(running) / len(pwm) * 100:.1f}%'))
        lines.append(('平均运行档位(仅运行时)', f'{sum(running) / len(running):.0f}（约 {sum(running) / len(running) / 255 * 100:.0f}%）' if running else '无'))
    lines.append(('', ''))

    ssd = numeric_column(rows, 'ssd_c')
    sources = sorted({row.get('ssd_source', '') for row in rows if row.get('ssd_source')})
    lines.append(('SSD 温度来源', ', '.join(sources) if sources else '(无)'))
    if ssd:
        lines.append(('SSD 温度(°C)', f'{min(ssd):.1f} / {sum(ssd) / len(ssd):.1f} / {max(ssd):.1f}'))
    else:
        lines.append(('SSD 温度(°C)', '无可用数据（权限或工具缺失）'))
    lines.append(('', ''))

    load1 = numeric_column(rows, 'load1')
    if load1:
        lines.append(('负载 load1', f'{min(load1):.2f} / {sum(load1) / len(load1):.2f} / {max(load1):.2f}'))
    mem = numeric_column(rows, 'mem_avail_mb')
    if mem:
        lines.append(('可用内存(MB)', f'{min(mem):.0f} / {sum(mem) / len(mem):.0f} / {max(mem):.0f}'))
    return lines


def add_chart(sheet, headers, row_count, title, columns, y_title, anchor):
    selected = [name for name in columns if name in headers]
    if not selected:
        return anchor
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = '采样序号'
    chart.height = 9
    chart.width = 28
    for name in selected:
        index = headers.index(name) + 1
        reference = Reference(sheet, min_col=index, min_row=1, max_row=row_count + 1)
        chart.add_data(reference, titles_from_data=True)
    sheet.add_chart(chart, anchor)
    row = int(''.join(character for character in anchor if character.isdigit()) or '2')
    return f'{get_column_letter(len(headers) + 2)}{row + 19}'


def main(argv):
    if not argv:
        print(__doc__)
        return 1
    source = argv[0]
    target = argv[1] if len(argv) > 1 else source.rsplit('.', 1)[0] + '.xlsx'
    rows = load_rows(source)
    if not rows:
        print('CSV 没有数据行')
        return 1
    headers = list(rows[0].keys())

    workbook = Workbook()
    data = workbook.active
    data.title = 'data'
    data.append(headers)
    for row in rows:
        data.append([
            row.get(name) if name in TEXT_COLUMNS else to_number(row.get(name))
            for name in headers
        ])
    for cell in data[1]:
        cell.font = Font(bold=True)
    data.freeze_panes = 'A2'

    anchor_row = 2
    for title, columns, unit in [
        ('SoC 各域温度', TEMPERATURE_COLUMNS, '°C'),
        ('风扇 PWM 与负载', ['pwm', 'load1', 'load5'], 'pwm / load'),
        ('CPU 频率', ['cpu0_mhz', 'cpu4_mhz', 'cpu6_mhz'], 'MHz'),
        ('SSD 温度', ['ssd_c'], '°C'),
    ]:
        anchor_row = add_chart(
            data, headers, len(rows), title, columns, unit,
            f'{get_column_letter(len(headers) + 2)}{anchor_row}',
        )

    report = workbook.create_sheet('report')
    report.append(['项目', '数值'])
    for cell in report[1]:
        cell.font = Font(bold=True)
    for name, value in summarize(rows):
        report.append([name, value])
    report.column_dimensions['A'].width = 26
    report.column_dimensions['B'].width = 40

    workbook.save(target)
    print(f'已生成 {target}：{len(rows)} 行数据，图表 4 张，摘要 {report.max_row - 1} 项')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
